# -*- coding: utf-8 -*-
import os
import re
from pathlib import Path
from datetime import datetime
import json
import traceback
from docx import Document
import firebase_admin
from firebase_admin import credentials, firestore
from dotenv import load_dotenv

# Base directories
BASE_DIR = Path(__file__).parent.parent
MD_DIR = BASE_DIR / 'md' / 'news'
NOTICIAS_DIR = BASE_DIR / 'noticias'
ASSETS_NEWS_DIR = BASE_DIR / 'assets' / 'news'
TEMPLATE_PATH = NOTICIAS_DIR / 'post_template.html'

# Ensure directories exist
NOTICIAS_DIR.mkdir(parents=True, exist_ok=True)
ASSETS_NEWS_DIR.mkdir(parents=True, exist_ok=True)

def slugify(text):
    text = text.lower()
    text = re.sub(r'[^\w\s-]', '', text)
    text = re.sub(r'[\s_-]+', '-', text).strip('-')
    return text

def initialize_firebase():
    """Initialize Firebase Admin SDK using environment variables from .env"""
    env_path = BASE_DIR / '.env'
    if not env_path.exists():
        print("ERROR: .env file not found at " + str(env_path))
        return None
    
    load_dotenv(env_path)
    
    try:
        # Check if already initialized
        try:
            return firestore.client()
        except ValueError:
            pass

        private_key = os.getenv("FIREBASE_PRIVATE_KEY", "").replace('\\n', '\n')
        cred_dict = {
            "type": os.getenv("FIREBASE_TYPE", "service_account"),
            "project_id": os.getenv("FIREBASE_PROJECT_ID"),
            "private_key_id": os.getenv("FIREBASE_PRIVATE_KEY_ID"),
            "private_key": private_key,
            "client_email": os.getenv("FIREBASE_CLIENT_EMAIL"),
            "client_id": os.getenv("FIREBASE_CLIENT_ID"),
            "auth_uri": os.getenv("FIREBASE_AUTH_URI"),
            "token_uri": os.getenv("FIREBASE_TOKEN_URI"),
            "auth_provider_x509_cert_url": os.getenv("FIREBASE_AUTH_PROVIDER_CERT_URL"),
            "client_x509_cert_url": os.getenv("FIREBASE_CLIENT_CERT_URL"),
            "universe_domain": os.getenv("FIREBASE_UNIVERSE_DOMAIN", "googleapis.com")
        }
        
        if not cred_dict["project_id"] or not cred_dict["private_key"]:
            print("ERROR: Missing required Firebase credentials")
            return None
            
        cred = credentials.Certificate(cred_dict)
        firebase_admin.initialize_app(cred)
        return firestore.client()
        
    except Exception as e:
        print("ERROR initializing Firebase: " + str(e))
        return None

def extract_images(doc, slug):
    """Extracts images from docx and saves to assets/news/"""
    images = []
    try:
        for i, rel in enumerate(doc.part.rels.values()):
            if "image" in rel.target_ref:
                img_data = rel.target_part.blob
                img_ext = Path(rel.target_ref).suffix
                img_filename = f"{slug}-{i}{img_ext}"
                img_path = ASSETS_NEWS_DIR / img_filename
                
                with open(img_path, "wb") as f:
                    f.write(img_data)
                
                images.append(f"assets/news/{img_filename}")
    except Exception as e:
        print(f"Error extracting images: {e}")
    return images

def process_docx(file_path):
    print(f"📄 Processing: {file_path.name}")
    doc = Document(file_path)
    
    title = ""
    content_html = ""
    paragraphs = []
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
            
        # First non-empty paragraph is the title if none found
        if not title:
            title = text
            continue
            
        paragraphs.append(text)
        content_html += f"<p>{text}</p>\n"
    
    if not title:
        title = file_path.stem
        
    slug = slugify(title)
    
    # Extract images
    images = extract_images(doc, slug)
    main_image = f"/{images[0]}" if images else "https://images.unsplash.com/photo-1574629810360-7efbbe195018?auto=format&fit=crop&q=80&w=2000"
    
    months = {
        1: "Janeiro", 2: "Fevereiro", 3: "Março", 4: "Abril",
        5: "Maio", 6: "Junho", 7: "Julho", 8: "Agosto",
        9: "Setembro", 10: "Outubro", 11: "Novembro", 12: "Dezembro"
    }
    now = datetime.now()
    date_str = f"{now.day} de {months[now.month]} de {now.year}"
    
    # Subtitle: use the second paragraph if available, otherwise the first
    subtitle = paragraphs[1] if len(paragraphs) > 1 else (paragraphs[0] if paragraphs else "")
    if len(subtitle) > 160:
        subtitle = subtitle[:157] + "..."
        
    return {
        "title": title,
        "subtitle": subtitle,
        "content_html": content_html,
        "slug": slug,
        "main_image": main_image,
        "date": date_str,
        "timestamp": now
    }

def generate_html_page(news_data):
    if not TEMPLATE_PATH.exists():
        print("❌ Template not found!")
        return
        
    with open(TEMPLATE_PATH, "r", encoding="utf-8") as f:
        html = f.read()
        
    # Replace placeholders
    html = html.replace("{{TITLE}}", news_data["title"])
    html = html.replace("{{SUBTITLE}}", news_data["subtitle"])
    html = html.replace("{{CONTENT}}", news_data["content_html"])
    html = html.replace("{{IMAGE_URL}}", news_data["main_image"])
    html = html.replace("{{DATE}}", news_data["date"])
    html = html.replace("{{CATEGORY}}", "Notícias") # Default
    
    output_path = NOTICIAS_DIR / f"{news_data['slug']}.html"
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html)
        
    print(f"✅ Generated HTML: {output_path.name}")
    return f"/noticias/{output_path.name}"

def publish_to_firestore(db, news_data, html_url):
    if not db:
        return
        
    try:
        news_ref = db.collection('news').document(news_data['slug'])
        
        # Check if already exists to avoid overwriting manually edited fields if desired
        # doc = news_ref.get()
        
        news_ref.set({
            "title": news_data["title"],
            "subtitle": news_data["subtitle"],
            "image_url": news_data["main_image"],
            "source_url": html_url,
            "last_updated_date_time": news_data["timestamp"],
            "is_highlight": False,
            "category": "Geral"
        })
        print(f"🔥 Published to Firestore: {news_data['title']}")
    except Exception as e:
        print(f"❌ Error publishing to Firestore: {e}")

def main():
    db = initialize_firebase()
    
    docx_files = list(MD_DIR.glob("*.docx"))
    if not docx_files:
        print("📭 No .docx files found in md/news/")
        return
        
    for file_path in docx_files:
        try:
            news_data = process_docx(file_path)
            html_url = generate_html_page(news_data)
            publish_to_firestore(db, news_data, html_url)
        except Exception as e:
            print(f"❌ Failed to process {file_path.name}: {e}")
            traceback.print_exc()

if __name__ == "__main__":
    main()
